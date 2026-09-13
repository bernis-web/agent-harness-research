# -*- coding: utf-8 -*-
"""在模板副本上填充科研实践计划书(保留学校模板格式)。"""
import copy
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

doc = Document(r'E:\科研实践2026秋\科研实践计划书模板(2).docx')


def set_cn_font(run, name='宋体', size=12, bold=False):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.bold = bold


def style_para(np, size, indent, font):
    np.paragraph_format.line_spacing = 1.3
    if indent:
        np.paragraph_format.first_line_indent = Pt(size * 2)
    else:
        np.paragraph_format.first_line_indent = Pt(0)


def fill_para(para, text, size=12, bold=False, font='宋体', align=None, indent=True):
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    run = para.add_run(text)
    set_cn_font(run, font, size, bold)
    if align is not None:
        para.alignment = align
    style_para(np=para, size=size, indent=indent, font=font)
    return para


def make_para(text, size=12, bold=False, font='宋体', indent=True):
    """构造一个独立的 w:p 元素(尚未插入文档)。"""
    p_elem = OxmlElement('w:p')
    np = Paragraph(p_elem, None)
    run = np.add_run(text)
    set_cn_font(run, font, size, bold)
    style_para(np, size, indent, font)
    return p_elem, np


def insert_after_elem(elem, text, size=12, bold=False, font='宋体', indent=True):
    p_elem, np = make_para(text, size, bold, font, indent)
    elem.addnext(p_elem)
    return np, p_elem


class Anchor:
    """跟踪"当前最后一个元素",支持段落与表格混排。"""

    def __init__(self, elem, parent):
        self.elem = elem
        self.parent = parent

    def add(self, text, size=12, bold=False, font='宋体', indent=True):
        p_elem, np = make_para(text, size, bold, font, indent)
        self.elem.addnext(p_elem)
        self.elem = p_elem
        return np

    def add_heading(self, text, size=13):
        return self.add(text, size=size, bold=True, font='黑体', indent=False)

    def add_table(self, table):
        self.elem.addnext(table._tbl)
        self.elem = table._tbl


# ---------- 定位模板锚点 ----------
paras = doc.paragraphs
p_title = paras[5]
p_course = paras[12]
p_meaning = paras[21]
p_status = paras[22]
p_content = paras[23]

# ---------- 封面 ----------
fill_para(p_title, '智能体Harness架构研究——以deepseek-harness与LangChain Deep Agents为对象',
          size=16, bold=True, font='黑体', align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
fill_para(p_course, '课程/项目名称：智能体Harness架构研究（本科生科研实践）', size=12, indent=False)

# ---------- 一、研究意义 ----------
a = Anchor(p_meaning._p, p_meaning._parent)
for t in [
    '2025年以来，智能体领域正在形成一个新的基本共识：Agent = 模型 + Harness。模型负责单步推理，而Harness（智能体运行时框架）负责让推理持续、可靠、可控地运转——对话循环、工具系统、权限管线、上下文管理与会话持久化都属于这一层。随着主流模型能力快速趋同，同样的模型接入不同运行时表现差异巨大，Harness层的工程设计成为智能体效果差异的主要来源。2026年7月底，DeepSeek随V4-Flash模型开源了deepseek-harness（dsh），上线12小时GitHub星标突破5万，把Agent Harness从一个工程术语推向了独立的技术品类。',
    'Harness的通用骨架恰好由六大模块构成：记忆管理、上下文管理、知识库管理、工具管理、复杂任务执行与多智能体协作。两个相互独立的开源项目（dsh与LangChain Deep Agents）不约而同地把运行时拆成这六个部分，说明研究这六块等于研究智能体运行时的"解剖学"。只有理解智能体的底层运行逻辑，才能设计和使用好智能体——这正是本课题的出发点，也与课程"追求一定科研程度、研究稍微底层内容"的要求一致。',
    '本课题的价值有三：其一，dsh与Deep Agents均为开源项目，可读源码、可复现、可改造，适合本科生科研实践的深度要求与实验条件；其二，Harness层的若干关键机制（如会话压缩的信息保真度、上下文管理与推理缓存命中的冲突）目前缺乏系统的定量研究，存在可做出增量贡献的空白；其三，研究成果（中文机制文档集、受控实验基准、可运行的扩展实现）对中文智能体开发社区有独立的参考价值。',
]:
    a.add(t)

# ---------- 二、研究现状及实验室研究基础 ----------
a = Anchor(p_status._p, p_status._parent)
a.add_heading('（一）研究现状')
for t in [
    'Harness概念已由Claude Code确立为范式，社区已有系统化解剖资料（如开源书《御舆——解码Agent Harness》）；学术界与之对应的是SWE-agent提出的ACI（Agent-Computer Interface）概念，证明"给模型什么接口"本身显著影响任务成功率。六个子方向的研究现状概述如下：',
    '记忆管理：MemGPT（Letta）提出操作系统式分层记忆，Mem0提供生产级记忆API，A-MEM（NeurIPS 2025）与RMM（ACL 2025）代表自主记忆架构的前沿；但"记忆以何种形式、何时注入上下文"这一Harness层决策缺乏受控对比。',
    '上下文管理：Anthropic提出上下文工程（compaction／工具结果清理／外部记忆三件套），Manus公开KV-cache友好设计经验；"上下文卫生与缓存命中的冲突"没有被定量刻画过。',
    '知识库管理：检索增强沿Naive→Advanced→GraphRAG→Agentic RAG演进；过程性知识由Agent Skills规范（SKILL.md渐进式披露）承接，两个研究对象均已支持。',
    '工具管理：MCP（Anthropic，2024.11）成为工具接入事实标准，但工具规模从十个涨到上百个时"如何把工具供给模型"（全量注入／分组／检索／两阶段）缺乏系统数据。',
    '复杂任务执行：从ReAct到"规划即工具"（write_todos、plan模式），长任务成功率依赖压缩、状态外置与子代理隔离的共同支撑，但压缩造成的信息损失无人测量。',
    '多智能体协作：应用内编排（LangGraph／AutoGen）与协议层（MCP／A2A）分层互补；Anthropic（多智能体更优）与Cognition（上下文分裂有害）的公开争议缺少小规模可复现的受控实证。',
]:
    a.add(t)
a.add_heading('（二）实验室研究基础')
a.add('本课题已在学期前段完成三部分基础工作。其一，完成了dsh源码的十篇中文机制解剖笔记（基于commit c291e79，v0.1.5-rc.2），覆盖仓库地图与包依赖、微内核与Cordis依赖注入、插件生命周期与装配、事件溯源会话、工具执行管线、goal/plan/todo任务层级、子代理与工作流、压缩与结果外置、模型接缝与MCP、声明式组合与扩展；每篇含行号级代码引用、机制图与设计权衡，并附可复核的校验命令，其中前三篇已经过约40项引用逐条复核的独立评审并修订定稿。其二，完成了领域调研报告，梳理六大子方向脉络并提出7个可切入的研究问题，形成7份可直接执行的候选实验方案。其三，搭建了可复现的实验环境：dsh与Deep Agents仓库本地克隆、Node.js与Python双栈、DeepSeek API接入与知识图谱工具链。上述工作全部开源，后续研究在既有基础上直接推进。')

# ---------- 三、研究内容及研究计划 ----------
a = Anchor(p_content._p, p_content._parent)
a.add_heading('（一）研究内容')
a.add('本课题按"机制解剖—受控实验—扩展实践"三个层次递进：')
a.add('（1）机制解剖收尾与交叉验证。以dsh为主线、Deep Agents为对照，完成六大模块的机制图谱（记忆、上下文、知识库、工具、任务、多智能体各对应什么运行时机制、如何衔接），完成十篇解剖笔记的收尾、交叉核验与合集整理，并产出一张总架构图与三框架（dsh／Deep Agents／Claude Code）对比表。')
a.add('（2）受控实验（核心科研增量）。设计并运行"长会话压缩保真度基准"：控制模型与任务不变，对比四种历史压缩策略（朴素截断／全文摘要／结构化摘要／文件外置+检索回填）在多约束长任务上的端到端成功率、关键约束召回率与token成本；任务集在早期埋入延迟使用的约束，使压缩损失可被末端评分捕获。在时间允许时追加"KV-cache友好的上下文管理"小实验：利用DeepSeek API返回的缓存命中计费字段，量化压缩与清理策略对推理缓存命中的影响。')
a.add('（3）扩展实践。基于实验结论，在Deep Agents上实现一个改进的压缩或记忆注入中间件（如结构化摘要策略），用同一基准验证其相对基线的增益，形成"研究—实现—验证"闭环。')
a.add_heading('（二）研究计划（16周）')

table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr = table.rows[0].cells
for i, h in enumerate(['周次', '任务', '阶段产出']):
    run = hdr[i].paragraphs[0].add_run(h)
    run.bold = True
for r in [
    ('1—3周', '精读两框架源码与核心文献（ReAct／MemGPT／SWE-agent／A-MEM等）；解剖笔记收尾', '机制文档集＋文献笔记'),
    ('4—6周', '实验设计：任务集生成器、自动评分器、四种压缩策略实现', '基准代码与任务集'),
    ('7—9周', '主实验（4策略×20任务×2重复）与消融（压缩触发阈值）', '实验数据与分析图表'),
    ('10—11周', 'KV-cache小实验，或改进型中间件的实现', '第二组数据／可运行扩展'),
    ('12—13周', '数据整理、结论提炼、失败案例分析', '实验报告初稿'),
    ('14—16周', '期末论文撰写、答辩准备、代码仓库整理开源', '论文＋答辩材料＋开源仓库'),
]:
    cells = table.add_row().cells
    for i, v in enumerate(r):
        cells[i].paragraphs[0].add_run(v)
a.add_table(table)

a.add_heading('（三）预期成果')
a.add('（1）智能体Harness机制的中文文档合集（十篇机制笔记＋总架构图＋三框架对比表），全部开源；（2）一组可复现的压缩保真度基准数据与结论，回答"不同压缩策略各丢失什么信息、成本与质量如何权衡"；（3）一个经过基准验证的可运行扩展（压缩或记忆注入中间件／dsh插件）；（4）期末论文与答辩材料。')
a.add_heading('（四）主要参考文献')
for t in [
    '[1] Yao et al. ReAct: Synergizing Reasoning and Acting in Language Models. ICLR 2023.',
    '[2] Packer et al. MemGPT: Towards LLMs as Operating Systems. 2023.',
    '[3] Yang et al. SWE-agent: Agent-Computer Interfaces Enable Automated Software Engineering. NeurIPS 2024.',
    '[4] Xu et al. A-MEM: Agentic Memory for LLM Agents. NeurIPS 2025.',
    '[5] Anthropic. Effective Context Engineering for AI Agents. 2025.',
    '[6] Manus. Context Engineering for AI Agents: Lessons from Building Manus. 2025.',
    '[7] 智能体AI框架：架构、协议与设计挑战. arXiv:2508.10146, 2025.',
    '[8] Model Context Protocol Specification. https://modelcontextprotocol.io.',
    '[9] deepseek-ai/deepseek-harness. https://github.com/deepseek-ai/deepseek-harness, 2026.',
    '[10] langchain-ai/deepagents. https://github.com/langchain-ai/deepagents.',
]:
    a.add(t, indent=False)

out = r'E:\科研实践2026秋\科研实践计划书-智能体Harness架构研究.docx'
doc.save(out)
print('saved:', out)
